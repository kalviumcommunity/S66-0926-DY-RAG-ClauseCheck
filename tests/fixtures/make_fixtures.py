"""Builds small, deterministic sample documents for tests (no binary files committed)."""
from pathlib import Path

import docx
import fitz  # PyMuPDF

PDF_PAGE_1 = (
    "SUPPLY AGREEMENT\n"
    "1. General provisions\n"
    "1.1 The Supplier shall deliver the Products described in Schedule A.\n"
    "1.2 This Agreement is governed by the laws of Delaware.\n"
    "Source: ACME CORP, 10-K, 1/15/2019"
)
PDF_PAGE_2 = (
    "2. Payment\n"
    "2.1 Payment is due Net 45 days from the invoice date.\n"
    "2.2 The rebate shall be [***] percent of annual spend.\n"
    "3. Term\n"
    "3.1 The initial term is two (2) years and renews automatically.\n"
    "Source: ACME CORP, 10-K, 1/15/2019"
)

HTML = """<html><head><style>p{}</style><script>var x=1;</script></head><body>
<p>AMENDMENT NO. 1 TO SUPPLY AGREEMENT</p>
<p>1. Section 2.1 of the Agreement is amended to read: Payment is due Net 30 days.</p>
<p>2. All other terms remain in effect.</p>
<table><tr><th>Item</th><th>Old</th><th>New</th></tr>
<tr><td>Payment terms</td><td>Net 45</td><td>Net 30</td></tr></table>
<p>Source: ACME CORP, 8-K, 3/1/2020</p>
</body></html>"""

TXT = (
    "CERTIFICATE OF REGISTRATION\n"
    "ISO 27001:2022\n"
    "Supplier: Acme Corp\n"
    "Certificate number: A-114\n"
    "Valid from 2025-04-01\n"
    "Valid until 2027-03-31\n"
)

DOCX_HEADER = ["SKU", "Description", "Unit price", "Currency", "Payment terms"]
DOCX_ROWS = [
    ["SKU-001", "Steel bracket", "1,250.00", "USD", "Net 45"],
    ["SKU-002", "Copper wire 10m", "88.10", "USD", "Net 45"],
    ["SKU-003", "Fastener kit", "7.20", "USD", "Net 45"],
]


def build_all(out: Path) -> dict[str, Path]:
    out.mkdir(parents=True, exist_ok=True)
    paths = {}

    pdf = fitz.open()
    for text in (PDF_PAGE_1, PDF_PAGE_2):
        page = pdf.new_page()
        page.insert_text((50, 72), text, fontsize=10)
    paths["pdf"] = out / "acme_supply.pdf"
    pdf.save(paths["pdf"])

    blank = fitz.open()
    blank.new_page()
    paths["blank_pdf"] = out / "no_text_layer.pdf"
    blank.save(paths["blank_pdf"])

    paths["html"] = out / "acme_amendment.html"
    paths["html"].write_text(HTML, encoding="utf-8")

    d = docx.Document()
    d.add_paragraph("PRICE SCHEDULE - ACME CORP")
    d.add_paragraph("Effective date: 2025-01-01")
    table = d.add_table(rows=1 + len(DOCX_ROWS), cols=len(DOCX_HEADER))
    for i, h in enumerate(DOCX_HEADER):
        table.rows[0].cells[i].text = h
    for r, row in enumerate(DOCX_ROWS, start=1):
        for c, value in enumerate(row):
            table.rows[r].cells[c].text = value
    paths["docx"] = out / "acme_prices.docx"
    d.save(paths["docx"])

    paths["txt"] = out / "acme_iso27001.txt"
    paths["txt"].write_text(TXT, encoding="utf-8")
    return paths


MANIFEST_ROWS = {
    "pdf": {"doc_id": "cuad-acme01", "supplier_id": "S-ACME", "source": "cuad",
            "doc_type": "contract", "format": "pdf", "title": "Acme Supply Agreement",
            "effective_date": "2019-01-15", "parent_doc_id": "", "is_synthetic": "false"},
    "html": {"doc_id": "edgar-acme02", "supplier_id": "S-ACME", "source": "edgar",
             "doc_type": "amendment", "format": "html", "title": "Amendment No. 1",
             "effective_date": "2020-03-01", "parent_doc_id": "cuad-acme01", "is_synthetic": "false"},
    "docx": {"doc_id": "synthetic-acme03", "supplier_id": "S-ACME", "source": "synthetic",
             "doc_type": "price_schedule", "format": "docx", "title": "Acme price schedule",
             "effective_date": "2025-01-01", "parent_doc_id": "", "is_synthetic": "true"},
    "txt": {"doc_id": "synthetic-acme04", "supplier_id": "S-ACME", "source": "synthetic",
            "doc_type": "certificate", "format": "txt", "title": "Acme ISO 27001",
            "effective_date": "2025-04-01", "parent_doc_id": "", "is_synthetic": "true"},
}
