import importlib.util
import json
from pathlib import Path

import docx

from clausecheck.core import schema

ROOT = Path(__file__).resolve().parents[2]


def _load():
    spec = importlib.util.spec_from_file_location("generate_synthetic", ROOT / "scripts/generate_synthetic.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _content(out: Path) -> dict[str, str]:
    """Text content of every generated file (DOCX zips carry timestamps, so compare text)."""
    result = {}
    for p in sorted(out.iterdir()):
        if p.suffix == ".docx":
            d = docx.Document(p)
            parts = [para.text for para in d.paragraphs]
            for t in d.tables:
                parts += ["|".join(c.text for c in row.cells) for row in t.rows]
            result[p.name] = "\n".join(parts)
        else:
            result[p.name] = p.read_text(encoding="utf-8")
    return result


def _strip(rows):
    return [{k: v for k, v in r.items() if k != "sha256"} for r in rows]


def test_same_seed_same_output(tmp_path):
    gen = _load()
    rows_a = gen.generate(42, tmp_path / "a")
    rows_b = gen.generate(42, tmp_path / "b")
    assert _strip(rows_a) == _strip(rows_b)
    assert _content(tmp_path / "a") == _content(tmp_path / "b")


def test_different_seed_differs(tmp_path):
    gen = _load()
    gen.generate(42, tmp_path / "a")
    gen.generate(7, tmp_path / "b")
    assert _content(tmp_path / "a") != _content(tmp_path / "b")


def test_rows_are_valid_manifest_rows(tmp_path):
    rows = _load().generate(42, tmp_path / "out")
    assert rows
    for r in rows:
        assert list(r.keys()) == schema.MANIFEST_COLUMNS
        assert r["source"] == "synthetic" and r["is_synthetic"] == "true"
        assert r["doc_type"] in {"price_schedule", "certificate"}
        assert r["format"] in {"docx", "txt"}


def test_price_schedules_have_one_table_with_header(tmp_path):
    out = tmp_path / "out"
    _load().generate(42, out)
    files = list(out.glob("price_schedule_*.docx"))
    assert files
    for p in files:
        tables = docx.Document(p).tables
        assert len(tables) == 1
        header = [c.text for c in tables[0].rows[0].cells]
        assert header == ["SKU", "Description", "Unit price", "Currency", "Payment terms"]
        assert 5 <= len(tables[0].rows) - 1 <= 15


def test_ground_truth_matches_rows(tmp_path):
    out = tmp_path / "out"
    rows = _load().generate(42, out)
    ids = {r["doc_id"] for r in rows}
    lines = [json.loads(l) for l in (out / "ground_truth.jsonl").read_text().splitlines() if l]
    assert lines
    assert all({"doc_id", "field", "value"} <= l.keys() for l in lines)
    assert {l["doc_id"] for l in lines} <= ids
