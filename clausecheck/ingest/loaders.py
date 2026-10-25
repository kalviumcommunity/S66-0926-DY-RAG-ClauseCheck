"""Turn a manifest row (PDF, HTML, DOCX or TXT) into a `RawDocument`."""

from pathlib import Path

import docx
import pdfplumber
import pymupdf
from bs4 import BeautifulSoup

from clausecheck.ingest.models import Page, RawDocument, Table

FORMAT_ALIASES = {"htm": "html"}
REMOVED_HTML_TAGS = ["script", "style", "nav", "header", "footer"]


class UnsupportedFormatError(ValueError):
    """Raised for a format the loaders do not handle."""


class EmptyDocumentError(ValueError):
    """Raised when a document has no text and no fallback text is available."""


def _cell(value: object) -> str:
    return " ".join(str(value or "").split())


def _table(page: int, rows: list[list[str]]) -> Table | None:
    """Build a Table from rows (first row = header); None if it has fewer than 2 rows."""
    if len(rows) < 2:
        return None
    return Table(page=page, header=rows[0], rows=rows[1:])


def _pages_from_chunks(chunks: list[str]) -> list[Page]:
    """Number the non-empty chunks as pages, or return one empty page if there are none."""
    texts = [c for c in (chunk.strip() for chunk in chunks) if c]
    return [Page(number=i, text=t) for i, t in enumerate(texts, 1)] or [Page(1, "")]


def _load_pdf(path: Path) -> tuple[list[Page], list[Table]]:
    pages: list[Page] = []
    with pymupdf.open(path) as pdf:
        for i, page in enumerate(pdf, 1):
            pages.append(Page(number=i, text=page.get_text("text", sort=True)))
    tables: list[Table] = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            for raw in page.extract_tables():
                table = _table(i, [[_cell(c) for c in row] for row in raw])
                if table:
                    tables.append(table)
    return pages, tables


def _load_html(path: Path) -> tuple[list[Page], list[Table]]:
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "lxml")
    for tag in soup(REMOVED_HTML_TAGS):
        tag.decompose()
    tables = []
    for element in soup.find_all("table"):
        rows = [
            [_cell(c.get_text(" ")) for c in tr.find_all(["th", "td"])]
            for tr in element.find_all("tr")
        ]
        table = _table(1, [r for r in rows if r])
        if table:
            tables.append(table)
    for hr in soup.find_all("hr"):
        hr.replace_with("\f")
    return _pages_from_chunks(soup.get_text("\n").split("\f")), tables


def _load_docx(path: Path) -> tuple[list[Page], list[Table]]:
    document = docx.Document(str(path))
    text = "\n".join(p.text for p in document.paragraphs)
    tables = []
    for element in document.tables:
        table = _table(1, [[_cell(c.text) for c in row.cells] for row in element.rows])
        if table:
            tables.append(table)
    return [Page(number=1, text=text)], tables


def _load_txt(path: Path) -> tuple[list[Page], list[Table]]:
    text = path.read_bytes().decode("utf-8", errors="replace")
    return _pages_from_chunks(text.split("\f")), []


_LOADERS = {"pdf": _load_pdf, "html": _load_html, "docx": _load_docx, "txt": _load_txt}


def load(path: Path, manifest_row: dict) -> RawDocument:
    """Load one document. The format comes from the manifest row, else the file suffix."""
    fmt = (manifest_row.get("format") or path.suffix.lstrip(".")).lower()
    fmt = FORMAT_ALIASES.get(fmt, fmt)
    if fmt not in _LOADERS:
        raise UnsupportedFormatError(f"unsupported format: {fmt!r} ({path.name})")
    pages, tables = _LOADERS[fmt](path)

    if not any(p.text.strip() for p in pages):
        fallback = manifest_row.get("fallback_text")
        if not fallback:
            raise EmptyDocumentError(f"no text in {path.name} and no fallback text")
        pages, tables = [Page(number=1, text=fallback)], []
    return RawDocument(
        doc_id=manifest_row["doc_id"],
        path=path,
        format=fmt,
        pages=pages,
        tables=tables,
        metadata=dict(manifest_row),
    )
